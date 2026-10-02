#!/usr/bin/env python3
from pathlib import Path
import sys, csv, json, zipfile, re


def extract_pdf(p):
    try:
        import fitz
        doc = fitz.open(str(p))
        parts = []
        for i, page in enumerate(doc):
            text = page.get_text("text").strip()
            if text:
                parts.append(f"\n### PDF Page {i+1}\n{text}")
        return "\n".join(parts)
    except Exception as e:
        return f"[PDF extraction failed: {e}]"


def extract_docx(p):
    try:
        from docx import Document
        doc = Document(str(p))
        out = [x.text for x in doc.paragraphs if x.text.strip()]
        for t in doc.tables:
            for row in t.rows:
                out.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(out)
    except Exception as e:
        return f"[DOCX extraction failed: {e}]"


def extract_pptx(p):
    try:
        from pptx import Presentation
        prs = Presentation(str(p))
        out = []
        for i, slide in enumerate(prs.slides, 1):
            out.append(f"\n### Slide {i}")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    out.append(shape.text.strip())
        return "\n".join(out)
    except Exception as e:
        return f"[PPTX extraction failed: {e}]"


def extract_xlsx(p):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(str(p), data_only=True, read_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"\n### Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = ["" if v is None else str(v) for v in row]
                if any(vals):
                    out.append(" | ".join(vals))
        return "\n".join(out)
    except Exception as e:
        return f"[XLSX extraction failed: {e}]"


def extract_csv_file(p):
    try:
        return p.read_text(encoding="utf-8-sig", errors="replace")
    except Exception as e:
        return f"[CSV extraction failed: {e}]"


def extract_text(p):
    return p.read_text(encoding="utf-8", errors="replace")


def main():
    if len(sys.argv) != 3:
        print("Usage: extract_sources.py <input_dir> <output_md>")
        sys.exit(2)
    input_dir = Path(sys.argv[1])
    output = Path(sys.argv[2])
    output.parent.mkdir(parents=True, exist_ok=True)

    files = sorted([p for p in input_dir.rglob("*") if p.is_file()])
    sections = ["# Extracted Weekly Case Sources\n"]
    if not files:
        sections.append("\n[No input files found.]\n")

    for p in files:
        ext = p.suffix.lower()
        rel = p.relative_to(input_dir)
        sections.append(f"\n\n# SOURCE: {rel}\n")
        if ext == ".pdf":
            body = extract_pdf(p)
        elif ext == ".docx":
            body = extract_docx(p)
        elif ext == ".pptx":
            body = extract_pptx(p)
        elif ext in {".xlsx", ".xlsm"}:
            body = extract_xlsx(p)
        elif ext == ".csv":
            body = extract_csv_file(p)
        elif ext in {".txt", ".md", ".text"}:
            body = extract_text(p)
        elif ext in {".png", ".jpg", ".jpeg", ".webp", ".gif"}:
            body = "[Image file: inspect visually with Claude Code image capabilities.]"
        else:
            body = f"[Unsupported by extractor: {ext}. Inspect with Claude Code if relevant.]"
        sections.append(body)

    output.write_text("\n".join(sections), encoding="utf-8")
    print(f"Wrote {output} from {len(files)} file(s).")


if __name__ == "__main__":
    main()
